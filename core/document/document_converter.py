import tempfile
from pathlib import Path
from typing import Callable, List, Optional
from core.conversion.base_converter import BaseConverter
from core.conversion.conversion_job import ConversionJob
from core.conversion.conversion_result import ConversionResult
from core.compression.compression_service import CompressionService
from core.exceptions import UnsupportedFormatError


class DocumentConverter(BaseConverter):
    """Conversor documental real con soporte cruzado entre PDF, DOCX y TXT."""

    def supported_source_formats(self) -> List[str]:
        return ["pdf", "doc", "docx", "txt"]

    def supported_target_formats(self) -> List[str]:
        return ["pdf", "docx", "txt"]

    def convert(
        self,
        job: ConversionJob,
        progress_callback: Optional[Callable[[int], None]] = None
    ) -> ConversionResult:
        try:
            src = job.source_format.lower()
            tgt = job.target_format.lower()
            temp_dir = Path(tempfile.gettempdir())
            temp_target = temp_dir / f"conver6_{job.id}.{tgt}"

            if src == "doc":
                return ConversionResult(
                    success=False,
                    error_message="El formato binario DOC requiere LibreOffice o Microsoft Word."
                )

            if progress_callback:
                progress_callback(20)

            if src == "txt" and tgt == "pdf":
                self._txt_to_pdf(job.input_path, temp_target, job.compression_level)
            elif src == "txt" and tgt == "docx":
                self._txt_to_docx(job.input_path, temp_target, job.compression_level)
            elif src == "pdf" and tgt == "txt":
                self._pdf_to_txt(job.input_path, temp_target)
            elif src == "pdf" and tgt == "docx":
                self._pdf_to_docx(job.input_path, temp_target)
            elif src == "docx" and tgt == "txt":
                self._docx_to_txt(job.input_path, temp_target)
            elif src == "docx" and tgt == "pdf":
                self._docx_to_pdf(job.input_path, temp_target)
            else:
                raise UnsupportedFormatError(f"Conversión no soportada: {src} -> {tgt}")

            if progress_callback:
                progress_callback(100)

            return ConversionResult(success=True, output_path=temp_target)

        except Exception as err:
            return ConversionResult(success=False, error_message=str(err))

    def _docx_to_pdf(self, src: Path, dest: Path) -> None:
        import docx
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT, TA_JUSTIFY

        doc_in = docx.Document(str(src))
        pdf = SimpleDocTemplate(
            str(dest),
            pagesize=letter,
            rightMargin=54,
            leftMargin=54,
            topMargin=54,
            bottomMargin=54
        )

        styles = getSampleStyleSheet()
        story = []

        align_map = {
            docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER: TA_CENTER,
            docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT: TA_RIGHT,
            docx.enum.text.WD_ALIGN_PARAGRAPH.JUSTIFY: TA_JUSTIFY,
        }

        for p in doc_in.paragraphs:
            text = p.text.strip()
            if not text:
                story.append(Spacer(1, 10))
                continue

            align = align_map.get(p.alignment, TA_LEFT)
            p_style = ParagraphStyle(
                name=f"Style_{len(story)}",
                parent=styles['Normal'],
                alignment=align,
                fontSize=11,
                leading=15,
                spaceAfter=6
            )
            # Reemplazar caracteres especiales HTML
            safe_text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(safe_text, p_style))

        if not story:
            story.append(Paragraph("", styles['Normal']))

        pdf.build(story)

    def _txt_to_pdf(self, src: Path, dest: Path, level: int) -> None:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

        text = src.read_text(encoding="utf-8", errors="replace")
        text = CompressionService.compress_text_content(text, level)
        pdf = SimpleDocTemplate(str(dest), pagesize=letter, leftMargin=54, rightMargin=54)
        styles = getSampleStyleSheet()
        style = ParagraphStyle('TXT', parent=styles['Normal'], fontSize=10, leading=14)
        story = [Paragraph(line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"), style)
                if line.strip() else Spacer(1, 8) for line in text.splitlines()]
        pdf.build(story or [Paragraph("", style)])

    def _txt_to_docx(self, src: Path, dest: Path, level: int) -> None:
        import docx
        doc = docx.Document()
        text = src.read_text(encoding="utf-8", errors="replace")
        text = CompressionService.compress_text_content(text, level)
        for line in text.splitlines():
            doc.add_paragraph(line)
        doc.save(str(dest))

    def _pdf_to_txt(self, src: Path, dest: Path) -> None:
        import pypdf
        reader = pypdf.PdfReader(str(src))
        content = [page.extract_text() or "" for page in reader.pages]
        dest.write_text("\n".join(content), encoding="utf-8")

    def _pdf_to_docx(self, src: Path, dest: Path) -> None:
        from pdf2docx import Converter
        cv = Converter(str(src))
        cv.convert(str(dest), start=0, end=None)
        cv.close()

    def _docx_to_txt(self, src: Path, dest: Path) -> None:
        import docx
        doc = docx.Document(str(src))
        content = [p.text for p in doc.paragraphs]
        dest.write_text("\n".join(content), encoding="utf-8")